#!/usr/bin/env python3
"""Migración única (oct 2026): mad-map-data-v2.xlsx -> investigacion-ead.xlsx.

Crea la planilla simple de siete hojas sin códigos internos y aplica la
redistribución de sublíneas acordada el 7 de octubre de 2026 (las sublíneas
cuya contribución es una herramienta, un servicio, un objeto o un sistema
pasan a Prácticas proyectuales, polo proyectual).
Se conserva sólo como registro; no volver a correr.
"""
import sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / (sys.argv[1] if len(sys.argv) > 1 else "mad-map-data-v2.xlsx")
DST = ROOT / "investigacion-ead.xlsx"
L1, L2 = "Fundamentos disciplinares", "Prácticas proyectuales"

A_PRACTICAS = {
    "Accesibilidad cognitiva y codiseño", "Accesibilidad e inclusión",
    "Comunicación aumentativa y alternativa",
    "Diseño de interacción, servicios y experiencia de usuario",
    "Diseño y vida independiente", "Diseño para la democracia y comunicación ciudadana",
    "Sistemas inteligentes e IA en diseño", "Máquinas expresivas, algoritmos y arte tecnológico",
    "Fabricación digital, modelado paramétrico y fablabs",
    "Industrias creativas e innovación social para el desarrollo local",
    "Transferencia de medios tecnológicos con aplicación al emprendimiento local",
    "Comunicación visual, diseño editorial y exposición material", "Saberes técnicos análogos",
}
NOTAS = {
    "Espacios de aprendizaje en contextos vulnerables": "Posible fusión con 'Espacios educativos en contextos vulnerables'.",
    "Espacios educativos en contextos vulnerables": "Posible fusión con 'Espacios de aprendizaje en contextos vulnerables'.",
    "Diseño arquitectónico": "Nombre vago; posible fusión con 'Teoría e historia de la arquitectura'.",
}
CLAUSTRO = {  # configuración A de propuesta-dos-lineas.md
    "Anna Braghini": L1, "Jaime Reyes Gil": L1, "Katherine Exss Cid": L1, "Arturo Chicano Jiménez": L1,  # Chicano: decisión 7 oct 2026
    "Álvaro Mercado Jara": L2, "Adriana Marín Toro": L2, "Emanuela Di Felice": L2, "Daniela Salgado Cofré": L2,
}
LINEAS_EXTRA = {
    L1: dict(prolonga="Educación, Espacio y Aprendizaje",
             pregunta="Qué tradición, qué oficio y qué teoría sostienen el proyecto de arquitectura y de diseño, cómo se actualizan sus categorías, y con qué espacios, métodos y medios se enseña y se aprende a proyectar.",
             labs="Patrimonio moderno; Archivo Histórico José Vial Armstrong"),
    L2: dict(prolonga="Extensión, Ciudad y Habitabilidad",
             pregunta="Cómo la ciudad y el territorio se construyen, se habitan, se sostienen y se piensan políticamente, y qué proyectos de arquitectura y de diseño (urbanos, de equipamiento, de objetos, servicios y sistemas) los transforman.",
             labs="Personas y territorios; Urbanismo afectivo; Núcleo de Accesibilidad e Inclusión; Aconcagua Fablab"),
}

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="B00000")


def rows(ws, hdr=4):
    h = [c.value for c in ws[hdr]]
    return [dict(zip(h, r)) for r in ws.iter_rows(min_row=hdr + 1, values_only=True) if r[0] not in (None, "")]


def sheet(wb, name, headers, widths):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for i, c in enumerate(ws[1]):
        c.font = HEAD; c.fill = FILL
    for col, w in zip("ABCDEFGH", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    return ws


def lista(ws, col, formula, n=600):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "Elige un valor del menú."; dv.showErrorMessage = True
    dv.add(f"{col}2:{col}{n}"); ws.add_data_validation(dv)


def main():
    src = load_workbook(SRC, data_only=True)
    lin = {r["nombre"]: r for r in rows(src["01_Lineas"])}
    subs = rows(src["02_Sublineas"])
    invs = rows(src["07_Investigadores"])
    temas = rows(src["08_Temas"])
    pry = rows(src["19_Proyectos"])
    sello = next(r for r in src["17_Sello"].iter_rows(values_only=True) if r[1] and "ELEGIDO" in str(r[1]))

    wb = Workbook(); wb.remove(wb.active)

    ws = sheet(wb, "Léeme", ["Cómo editar esta planilla"], [110])
    for t in [
        "Esta planilla es la única fuente de datos del sitio. Al guardar, hacer commit y push, el sitio se actualiza solo.",
        "",
        "Programa: la pregunta del programa y el sello. Una fila por campo; editar sólo la columna 'valor'.",
        "Líneas: las dos líneas. 'definición breve' es lo que se ve en la cabecera de cada columna del sitio.",
        "Sublíneas: una fila por sublínea. Para mover una sublínea, cambiar 'línea' y 'polo' con los menús.",
        "Profesores: una fila por profesor. 'claustro' indica la línea a la que pertenece en el claustro (vacío si no es parte).",
        "Temas: une a un profesor con una sublínea. Una fila por vínculo; 'tema' es opcional y describe lo que investiga ahí.",
        "Proyectos: una fila por proyecto. 'profesores' lleva los nombres separados por punto y coma (;), escritos igual que en la hoja Profesores.",
        "",
        "Reglas: los nombres deben escribirse siempre igual en todas las hojas (usar los menús cuando existan).",
        "Para agregar una sublínea o un profesor, agregar una fila al final de su hoja; aparecerá en los menús de las otras hojas.",
        "La línea y el polo de un proyecto no se escriben: se deducen de su sublínea.",
        "",
        "Pendientes de revisión: las sublíneas con nota y los proyectos marcados 'revisar = sí'.",
    ]:
        ws.append([t])
    for c in ws["A"]:
        c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = sheet(wb, "Programa", ["campo", "valor"], [22, 110])
    ws.append(["pregunta", "¿Cómo renovar, cada vez, el habitar humano?"])
    ws.append(["sello (título)", sello[1].replace(" (ELEGIDO)", "")])
    ws.append(["sello (texto)", sello[2]])
    for c in ws["B"]:
        c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = sheet(wb, "Líneas", ["nombre", "modo", "definición breve", "pregunta", "prolonga (área del Magíster)", "laboratorios", "descripción"], [26, 34, 60, 60, 32, 40, 80])
    for n in (L1, L2):
        r = lin[n]; e = LINEAS_EXTRA[n]
        ws.append([n, r.get("modo"), r.get("bajada"), e["pregunta"], e["prolonga"], e["labs"], r.get("descripción")])

    ws = sheet(wb, "Sublíneas", ["nombre", "línea", "polo", "notas"], [62, 26, 12, 60])
    movidas = 0
    for s in sorted(subs, key=lambda s: s["nombre"]):
        linea, polo = s["línea"], s["polo"]
        if s["nombre"] in A_PRACTICAS:
            linea, polo = L2, "proyectual"; movidas += 1
        ws.append([s["nombre"], linea, polo, NOTAS.get(s["nombre"])])
    lista(ws, "B", f'"{L1},{L2}"'); lista(ws, "C", '"teórico,proyectual"')

    ws = sheet(wb, "Profesores", ["nombre", "claustro", "perfil"], [32, 26, 60])
    for i in sorted(invs, key=lambda i: i["nombre"]):
        ws.append([i["nombre"], CLAUSTRO.get(i["nombre"]), i.get("perfil_casiopea")])
    lista(ws, "B", f'"{L1},{L2}"')

    ws = sheet(wb, "Temas", ["profesor", "sublínea", "tema"], [32, 62, 70])
    for t in sorted(temas, key=lambda t: (t["investigador"], t["sublínea"])):
        ws.append([t["investigador"], t["sublínea"], t["tema"]])
    lista(ws, "A", "=Profesores!$A$2:$A$300"); lista(ws, "B", "='Sublíneas'!$A$2:$A$300")

    ws = sheet(wb, "Proyectos", ["año", "título", "profesores", "sublínea", "revisar"], [7, 80, 40, 50, 9])
    for p in sorted(pry, key=lambda p: (-(p["año"] or 0), p["título"])):
        ws.append([p["año"], p["título"], (p["investigadores"] or "").replace(", ", "; "), p["sublínea"], p["revisar"]])
    lista(ws, "D", "='Sublíneas'!$A$2:$A$300", 1000); lista(ws, "E", '"sí,no"', 1000)

    wb.save(DST)
    print(f"OK {DST.name}: {len(subs)} sublíneas ({movidas} movidas a {L2}), {len(invs)} profesores, {len(temas)} temas, {len(pry)} proyectos")


if __name__ == "__main__":
    main()
