#!/usr/bin/env python3
"""
Migra mad-map-data-v2.xlsx del régimen de cuatro líneas troncales al de dos
líneas de investigación (septiembre 2026):

    LIN-01  Formación, oficio y teoría del proyecto      (investigación acerca del proyecto)
    LIN-02  Proyecto, ciudad y ecologías del territorio  (investigación a través del proyecto)

Qué toca y qué no:
  - 01_Lineas: reemplaza las cuatro filas por dos.
  - 02_Sublineas: reasigna la columna `línea` por nombre y agrega la columna
    `polo` (teórico | proyectual), que expresa la contraparte interna de cada
    línea. Ninguna sublínea se elimina.
  - 10_Lab_Linea y 14_Linea_Modo: reescribe las relaciones por nombre.
  - 00_Lectura y 17_Sello: actualiza el texto que hablaba de cuatro líneas.
  - No toca temas, investigadores, laboratorios, modos ni proximidad.

Se ejecuta una sola vez desde la raíz del repositorio:

    python3 tools/migrate_dos_lineas.py

Deja una copia de seguridad del archivo original en
mad-map-data-v2.4lineas.bak.xlsx. Es idempotente: si ya se migró, no cambia nada.
Después conviene correr tools/build_doc.py para regenerar lineas-investigacion.md.
"""
import shutil
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "mad-map-data-v2.xlsx"
BACKUP = ROOT / "mad-map-data-v2.4lineas.bak.xlsx"

L1 = "Formación, oficio y teoría del proyecto"
L2 = "Proyecto, ciudad y ecologías del territorio"

LINEAS = [
    ("LIN-01", L1,
     "Investigación acerca del proyecto de arquitectura y de diseño: cómo se forma, cómo "
     "se ejerce como oficio y desde qué teoría e historia se piensa. Acervo de la e[ad] y "
     "Ciudad Abierta, historia y crítica de la arquitectura y del diseño, teoría del proyecto "
     "y sus categorías (hospitalidad, vacío, palabra poética), espacios y medios del "
     "aprendizaje, métodos de diseño, interacción, accesibilidad, fabricación y tecnologías. "
     "Forma a quienes enseñarán arquitectura y diseño. Polo teórico e histórico y polo "
     "proyectual.",
     "propuesta 2026"),
    ("LIN-02", L2,
     "Investigación a través del proyecto de arquitectura y de diseño sobre la ciudad, el "
     "territorio y sus ecologías: urbanización, ecología política, vivienda, patrimonio, "
     "borde costero, infraestructura y equipamiento, prácticas colectivas, urbanismo "
     "afectivo, diseño social y territorial. Polo teórico (teoría urbana, ecología "
     "política, vivienda) y polo proyectual (intervención, equipamiento, "
     "investigación-acción, creación de obra).",
     "propuesta 2026"),
]

# Sublínea (por id) -> (línea, polo). Todo lo que no aparece aquí cae en LIN-01
# proyectual, que es el destino de las antiguas LIN-01 (interacción) y LIN-04 (oficio).
ASIGNACION = {
    # Línea 1, polo teórico-histórico (antigua LIN-03)
    "SUB-03": (L1, "teórico"), "SUB-08": (L1, "teórico"), "SUB-12": (L1, "teórico"),
    "SUB-13": (L1, "teórico"), "SUB-26": (L1, "teórico"), "SUB-47": (L1, "teórico"),
    "SUB-48": (L1, "teórico"), "SUB-49": (L1, "teórico"), "SUB-50": (L1, "teórico"),
    "SUB-51": (L1, "teórico"), "SUB-52": (L1, "teórico"),
    # Línea 2, polo teórico
    "SUB-21": (L2, "teórico"), "SUB-29": (L2, "teórico"), "SUB-31": (L2, "teórico"),
    "SUB-32": (L2, "teórico"), "SUB-39": (L2, "teórico"),
    # Línea 2, polo proyectual
    "SUB-04": (L2, "proyectual"), "SUB-05": (L2, "proyectual"), "SUB-22": (L2, "proyectual"),
    "SUB-28": (L2, "proyectual"), "SUB-30": (L2, "proyectual"), "SUB-33": (L2, "proyectual"),
    "SUB-34": (L2, "proyectual"), "SUB-35": (L2, "proyectual"), "SUB-36": (L2, "proyectual"),
    "SUB-37": (L2, "proyectual"), "SUB-38": (L2, "proyectual"), "SUB-40": (L2, "proyectual"),
    "SUB-44": (L2, "proyectual"), "SUB-45": (L2, "proyectual"),
}
DEFAULT = (L1, "proyectual")

LAB_LINEA = [
    ("Núcleo de Accesibilidad e Inclusión", L1),
    ("Aconcagua Fablab", L1),
    ("Patrimonio moderno", L1),
    ("Personas y territorios", L2),
    ("Urbanismo afectivo", L2),
]

LINEA_MODO = [
    (L1, "Historiografía", "predominante"),
    (L1, "Teoría crítica", "predominante"),
    (L1, "Investigación proyectual", "predominante"),
    (L2, "Teoría crítica", "predominante"),
    (L2, "Investigación proyectual", "predominante"),
    (L2, "Historiografía", "presente"),
]

SELLO = (
    "El doctorado forma investigadores para quienes la obra es origen y prueba de la "
    "tesis. Esa obra, sea edificada, fabricada, escrita o dibujada, vale como argumento "
    "por sí misma, y el discurso doctoral se construye para hacerla legible y discutible. "
    "La pregunta que esa obra está llamada a argumentar es la que el programa hace suya: "
    "cómo reinventar el habitar humano, y la responde desde el habitar poético que funda "
    "a la Escuela. Las dos líneas del doctorado investigan ese habitar en arquitectura y "
    "en diseño: una investiga acerca del proyecto, su formación, su oficio y su teoría; la "
    "otra investiga a través del proyecto, sobre la ciudad, el territorio y sus ecologías. "
    "Cada línea lleva dentro su contraparte, teórica o proyectual. Doctorarse aquí es haber producido una obra capaz "
    "de sostener esos principios al ponerse a prueba y, en su mejor versión, capaz de "
    "reformularlos."
)


def header_row(ws, key):
    """Devuelve (índice de fila, {nombre de columna: índice}) de la fila de cabecera,
    identificada porque contiene la celda `key`. Se usa en cada hoja porque todas
    tienen una o dos filas de título antes de la cabecera real."""
    for row in ws.iter_rows(min_row=1, max_row=6):
        vals = [c.value for c in row]
        if key in vals:
            return row[0].row, {v: i + 1 for i, v in enumerate(vals) if v is not None}
    raise KeyError(f"cabecera con '{key}' no encontrada en {ws.title}")


def clear_below(ws, first_row, ncols):
    """Vacía las celdas de datos (sin borrar filas, para conservar formato y validaciones).
    Se usa antes de reescribir las hojas de relación."""
    for r in range(first_row, ws.max_row + 1):
        for c in range(1, ncols + 1):
            ws.cell(r, c).value = None


def migrar_lineas(wb):
    """Reemplaza las cuatro líneas por las dos nuevas en 01_Lineas."""
    ws = wb["01_Lineas"]
    hr, cols = header_row(ws, "id")
    clear_below(ws, hr + 1, 4)
    for i, (lid, nombre, desc, estado) in enumerate(LINEAS):
        r = hr + 1 + i
        ws.cell(r, cols["id"]).value = lid
        ws.cell(r, cols["nombre"]).value = nombre
        ws.cell(r, cols["descripción"]).value = desc
        ws.cell(r, cols["estado"]).value = estado
    ws.cell(1, 1).value = "01 · Líneas de investigación (2, propuesta septiembre 2026)"
    ws.cell(2, 1).value = ("Cada Sublínea pertenece a una Línea y a un polo (teórico | proyectual). "
                           "Los Laboratorios pueden sostener una o varias Líneas.")


def migrar_sublineas(wb):
    """Reasigna la columna `línea` y agrega/actualiza la columna `polo` en 02_Sublineas."""
    ws = wb["02_Sublineas"]
    hr, cols = header_row(ws, "id")
    if "polo" not in cols:
        cpolo = max(cols.values()) + 1
        ws.cell(hr, cpolo).value = "polo"
        cols["polo"] = cpolo
    for r in range(hr + 1, ws.max_row + 1):
        sid = ws.cell(r, cols["id"]).value
        if not sid:
            continue
        linea, polo = ASIGNACION.get(sid, DEFAULT)
        ws.cell(r, cols["línea"]).value = linea
        ws.cell(r, cols["polo"]).value = polo
    ws.cell(2, 1).value = ("Cada sublínea pertenece a UNA línea y a UN polo. Reasignación de "
                           "septiembre 2026 (dos líneas); revisar caso por caso.")


def migrar_relacion(wb, hoja, key, filas):
    """Reescribe una hoja de relación (10_Lab_Linea o 14_Linea_Modo) con las tuplas dadas."""
    ws = wb[hoja]
    hr, cols = header_row(ws, key)
    clear_below(ws, hr + 1, len(cols))
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila):
            ws.cell(hr + 1 + i, j + 1).value = v


def migrar_textos(wb):
    """Actualiza los textos de lectura y el sello que hablaban de cuatro líneas."""
    ws = wb["00_Lectura"]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("Línea (4 troncales)"):
                c.value = "Línea (2)"
            elif isinstance(c.value, str) and "Líneas troncales sostenidas" in c.value:
                c.value = ("Régimen de dos líneas de investigación (septiembre 2026), cada una con "
                           "polo teórico y polo proyectual, sostenidas por Laboratorios + sublíneas + temas.")
    ws = wb["17_Sello"]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "ELEGIDO" in c.value and c.column == 2:
                # La celda del texto puede ser parte de un rango combinado: se
                # escribe en la esquina superior izquierda de ese rango.
                target = ws.cell(c.row, c.column + 1)
                for rng in ws.merged_cells.ranges:
                    if target.coordinate in rng:
                        target = ws.cell(rng.min_row, rng.min_col)
                        break
                target.value = SELLO


def main():
    if not BACKUP.exists():
        shutil.copy2(XLSX, BACKUP)
        print(f"respaldo: {BACKUP.name}")
    wb = load_workbook(XLSX)
    migrar_lineas(wb)
    migrar_sublineas(wb)
    migrar_relacion(wb, "10_Lab_Linea", "laboratorio", LAB_LINEA)
    migrar_relacion(wb, "14_Linea_Modo", "línea", LINEA_MODO)
    migrar_textos(wb)
    wb.save(XLSX)
    print(f"OK: {XLSX.name} migrado a dos líneas")


if __name__ == "__main__":
    main()
