"""Lectura de investigacion-ead.xlsx para los scripts de documentos.

Mismo modelo que app.js: líneas, sublíneas (con línea y polo), profesores
(con línea y claustro), temas (profesor ↔ sublínea) y proyectos (con
sublínea; línea y polo se deducen de ella).
"""
import csv
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "investigacion-ead.xlsx"
PRODUCTIVIDAD = ROOT / "data" / "productividad.csv"


def _hoja(wb, nombre):
    ws = wb[nombre]
    filas = list(ws.iter_rows(values_only=True))
    h = [str(x).strip() if x else "" for x in filas[0]]
    out = []
    for r in filas[1:]:
        if r and r[0] not in (None, ""):
            out.append({k: (str(v).strip() if isinstance(v, str) else v) for k, v in zip(h, r) if k})
    return out


def cargar():
    wb = load_workbook(XLSX, data_only=True)
    programa = {r["campo"]: r["valor"] for r in _hoja(wb, "Programa")}
    lineas = _hoja(wb, "Líneas")
    subs = _hoja(wb, "Sublíneas")
    for s in subs:
        s["polo"] = "teórico" if str(s.get("polo", "")).lower().startswith("te") else "proyectual"
    sub_by = {s["nombre"]: s for s in subs}
    profs = _hoja(wb, "Profesores")
    for p in profs:
        p["claustro"] = str(p.get("claustro") or "").lower().startswith("s")
    temas = [t for t in _hoja(wb, "Temas") if t.get("sublínea") in sub_by]
    proyectos = []
    for p in _hoja(wb, "Proyectos"):
        s = sub_by.get(p.get("sublínea"))
        p["profesores"] = [x.strip() for x in str(p.get("profesores") or "").replace(",", ";").split(";") if x.strip()]
        p["línea"] = s["línea"] if s else None
        p["polo"] = s["polo"] if s else None
        p["revisar"] = str(p.get("revisar") or "").lower().startswith("s")
        proyectos.append(p)
    profs_sub = defaultdict(dict)  # sublínea -> {profesor: [temas]}
    for t in temas:
        profs_sub[t["sublínea"]].setdefault(t["profesor"], [])
        if t.get("tema"):
            profs_sub[t["sublínea"]][t["profesor"]].append(t["tema"])
    prod = {}
    if PRODUCTIVIDAD.exists():
        with open(PRODUCTIVIDAD) as f:
            for r in csv.DictReader(f):
                prod[r["investigador"]] = {k: int(v or 0) for k, v in r.items() if k != "investigador"}
    return dict(programa=programa, lineas=lineas, subs=subs, sub_by=sub_by, profs=profs,
                temas=temas, proyectos=proyectos, profs_sub=profs_sub, prod=prod)
