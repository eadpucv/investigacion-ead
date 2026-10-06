#!/usr/bin/env python3
"""Genera comparacion-lineas.md: comparación de las dos líneas del doctorado.

Combina texto editorial (definiciones, formas de investigar, contribución,
claustro) mantenido aquí con datos calculados desde mad-map-data-v2.xlsx:
sublíneas por polo (02_Sublineas), profesores por sublínea (08_Temas),
proyectos (19_Proyectos) y productividad (20_Productividad).

Uso:  python3 tools/build_comparacion.py
"""
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "mad-map-data-v2.xlsx"
OUT = ROOT / "comparacion-lineas.md"
L1, L2 = "Fundamentos disciplinares", "Prácticas proyectuales"

# --- Texto editorial por fila de la tabla comparativa -----------------------
EDIT = [
    ("Relación con el proyecto", "Investiga *acerca del* proyecto: el proyecto es su **objeto**", "Investiga *a través del* proyecto: el proyecto es su **método**"),
    ("Qué engloba", "Lo que sostiene a la disciplina: historia, teoría, filosofía y métodos, y cómo se transmiten", "Todo lo que se proyecta y se pone a prueba: obra nueva (un edificio, un espacio público, un objeto, un servicio, un sistema) que se inscribe en un contexto y lo modifica"),
    ("Pregunta nuclear", "Qué tradición, oficio y teoría sostienen el proyecto de arquitectura y de diseño, cómo se actualizan sus categorías, y con qué espacios, métodos y medios se enseña a proyectar", "Cómo la ciudad y el territorio se construyen, se habitan, se sostienen y se piensan políticamente, y qué proyectos de arquitectura y de diseño los transforman"),
    ("Prolonga (área del Magíster)", "Educación, Espacio y Aprendizaje", "Extensión, Ciudad y Habitabilidad"),
    ("Focos del polo teórico", "Acervo de la e[ad] y Ciudad Abierta; historia y crítica de la arquitectura moderna y latinoamericana; teoría del proyecto y sus categorías (hospitalidad, vacío, palabra poética); poesía y oficio; techné; reforma escolar", "Teoría urbana; urbanización y ecología política; vivienda y financiarización; comunes y resiliencias socioecológicas; perspectivas decoloniales; evaluación de políticas públicas"),
    ("Focos del polo proyectual", "Espacios educativos y arquitectura como medio didáctico; métodos de diseño; interacción, accesibilidad y comunicación aumentativa; sistemas inteligentes; fabricación digital; comunicación visual", "Urbanismo afectivo y deriva; investigación-acción; movilidad, infraestructura y equipamiento; ciudad-teatro; patrimonio y paisaje; diseño social y territorial; travesías y geopoética"),
    ("Formas de investigar", "Historiografía (archivo, fuentes, genealogías de obra); teoría crítica; investigación proyectual en el polo de métodos y aprendizaje", "Investigación proyectual; investigación-acción; creación de obra; teoría crítica con trabajo de campo; investigación especulativa"),
    ("Tipo de contribución (perfil de tesis)", "Tesis de historia, crítica o teoría del proyecto a partir de archivos y obras; exégesis de un dispositivo, espacio o medio de aprendizaje puesto a prueba; investigación sobre pedagogía del proyecto con obra docente como evidencia", "Tesis de teoría urbana, ecología política o vivienda con trabajo de campo; exégesis de una intervención urbana, un equipamiento, un objeto o un servicio; investigación-acción con comunidades e incidencia en política pública"),
    ("Competencias con más peso", "C1 (posicionamiento frente a una tradición) y C3 (conceptualización del aporte)", "C4 (configuración de la obra) y C5 (rendición de cuenta ante quienes la reciben)"),
    ("Objetivos del programa que sirve", "Énfasis docente y formativo del objetivo general; posicionamiento disciplinar (objetivo 1)", "Articulación con comunidades, territorios e instituciones (objetivo 3)"),
    ("Claustro (configuración A, recomendada)", "Polo teórico: Anna Braghini, Jaime Reyes. Polo proyectual: Katherine Exss", "Polo teórico: Álvaro Mercado, Adriana Marín. Polo proyectual: Emanuela Di Felice, Daniela Salgado"),
    ("Incorporaciones próximas al claustro", "Ursula Exss y Arturo Chicano (a un capítulo del umbral); Herbert Spencer al obtener el grado", "Iván Ivelic y Andrés Garcés (a un capítulo del umbral)"),
    ("Laboratorios", "Patrimonio moderno; Núcleo de Accesibilidad e Inclusión; Aconcagua Fablab; Archivo Histórico José Vial", "Personas y territorios; Urbanismo afectivo"),
    ("Interlocutores externos", "Escuelas de arquitectura y diseño, redes de historia y teoría, archivos", "Estado, gobiernos regionales, municipios, organizaciones de la sociedad civil, comunidades"),
    ("Fortaleza", "Nicho propio del programa: forma explícitamente a quienes enseñarán arquitectura y diseño", "El bloque más cohesionado del claustro y el de mayor financiamiento externo"),
    ("Riesgo o brecha", "Polo proyectual sostenido en el claustro sólo por K. Exss; núcleo histórico incompleto sin U. Exss", "Planificación urbana sin académico elegible (Baeriswyl es externo); Mercado y Salgado dejan de cumplir en 2025-2029 sin publicaciones nuevas"),
]
CLAUSTRO = {L1: ["Anna Braghini", "Jaime Reyes Gil", "Katherine Exss Cid"],
            L2: ["Álvaro Mercado Jara", "Adriana Marín Toro", "Emanuela Di Felice", "Daniela Salgado Cofré"]}


def rows(ws, hdr=4):
    h = [c.value for c in ws[hdr]]
    out = []
    for r in ws.iter_rows(min_row=hdr + 1, values_only=True):
        if r[0] not in (None, ""):
            out.append(dict(zip(h, r)))
    return out


def main():
    wb = load_workbook(XLSX, data_only=True)
    subs = rows(wb["02_Sublineas"])
    temas = rows(wb["08_Temas"])
    pry = rows(wb["19_Proyectos"])
    prod = {r["investigador"]: r for r in rows(wb["20_Productividad"])}

    sub_by_name = {s["nombre"]: s for s in subs}
    profs_sub = defaultdict(set)
    for t in temas:
        if t["sublínea"] in sub_by_name:
            profs_sub[t["sublínea"]].add(t["investigador"])
    profs_lin = defaultdict(set)
    for s, ps in profs_sub.items():
        profs_lin[sub_by_name[s]["línea"]] |= ps
    pry_sub = defaultdict(list)
    for p in pry:
        pry_sub[p["sublínea"]].append(p)

    def stats(lin):
        S = [s for s in subs if s["línea"] == lin]
        P = [p for p in pry if p["línea"] == lin]
        return dict(
            sub=len(S), subt=sum(s["polo"] == "teórico" for s in S), subp=sum(s["polo"] == "proyectual" for s in S),
            prof=len(profs_lin[lin]), pry=len(P), pry5=sum(p["ventana_2022_2026"] == "sí" for p in P),
            pryt=sum(p["polo"] == "teórico" for p in P), pryp=sum(p["polo"] == "proyectual" for p in P),
            activas=sum(1 for s in S if pry_sub[s["nombre"]]),
            cl_pub=sum(prod.get(n, {}).get("publicaciones_2022_2026", 0) or 0 for n in CLAUSTRO[lin]),
            cl_pry=sum(prod.get(n, {}).get("proyectos_2022_2026", 0) or 0 for n in CLAUSTRO[lin]),
        )
    A, B = stats(L1), stats(L2)

    o = []
    o.append("# Comparación de las líneas de investigación\n")
    o.append("**Doctorado en Arquitectura y Diseño**")
    o.append("*Escuela de Arquitectura y Diseño · Pontificia Universidad Católica de Valparaíso*\n")
    o.append(f"Las dos líneas del doctorado son simétricas: **{L1}** investiga *acerca del* proyecto y **{L2}** investiga *a través del* proyecto. "
             "Cada una lleva dentro su contraparte, un polo teórico y un polo proyectual. Este documento las compara en su definición, "
             "sus sublíneas, sus proyectos y su cuerpo académico.\n")
    o.append("> Documento generado por `tools/build_comparacion.py` desde `mad-map-data-v2.xlsx`. Los proyectos provienen del levantamiento de "
             "productividad 2014-2026 y se asignaron a una sublínea a partir de su título; las asignaciones de confianza media o baja "
             "están marcadas para revisión en la hoja `19_Proyectos`.\n")

    o.append("## Tabla comparativa\n")
    o.append(f"| | **{L1}** | **{L2}** |\n|---|---|---|")
    for k, a, b in EDIT[:8]:
        o.append(f"| **{k}** | {a} | {b} |")
    o.append(f"| **Sublíneas** | {A['sub']} ({A['subt']} teóricas, {A['subp']} proyectuales) | {B['sub']} ({B['subt']} teóricas, {B['subp']} proyectuales) |")
    o.append(f"| **Sublíneas con proyectos** | {A['activas']} de {A['sub']} | {B['activas']} de {B['sub']} |")
    o.append(f"| **Proyectos 2014-2026** | {A['pry']} ({A['pryt']} en el polo teórico, {A['pryp']} en el proyectual) | {B['pry']} ({B['pryt']} en el polo teórico, {B['pryp']} en el proyectual) |")
    o.append(f"| **Proyectos 2022-2026** | {A['pry5']} | {B['pry5']} |")
    o.append(f"| **Profesores que la cultivan** | {A['prof']} | {B['prof']} |")
    for k, a, b in EDIT[8:12]:
        o.append(f"| **{k}** | {a} | {b} |")
    o.append(f"| **Productividad del claustro 2022-2026** | {A['cl_pub']} publicaciones y {A['cl_pry']} proyectos ({len(CLAUSTRO[L1])} académicos) | {B['cl_pub']} publicaciones y {B['cl_pry']} proyectos ({len(CLAUSTRO[L2])} académicos) |")
    for k, a, b in EDIT[12:]:
        o.append(f"| **{k}** | {a} | {b} |")
    o.append("")
    o.append("Un mismo profesor puede cultivar sublíneas de ambas líneas, por lo que los conteos de profesores se superponen. En el claustro formal "
             "cada académico pertenece a una sola línea. La productividad del claustro cuenta publicaciones de cualquier tipo; el cotejo "
             "contra los criterios del área (WoS/Scopus o capítulos con referato) está en `propuesta-dos-lineas.md`.\n")

    for lin in (L1, L2):
        o.append(f"## {lin}: sublíneas\n")
        for polo in ("teórico", "proyectual"):
            o.append(f"### Polo {polo}\n")
            o.append("| Sublínea | Profesores | Proyectos | 2022-2026 |\n|---|---|---:|---:|")
            S = sorted([s for s in subs if s["línea"] == lin and s["polo"] == polo], key=lambda s: -len(pry_sub[s["nombre"]]))
            for s in S:
                ps = ", ".join(sorted(profs_sub[s["nombre"]])) or "—"
                P = pry_sub[s["nombre"]]
                o.append(f"| {s['nombre']} | {ps} | {len(P)} | {sum(p['ventana_2022_2026']=='sí' for p in P)} |")
            o.append("")

    for lin in (L1, L2):
        o.append(f"## {lin}: proyectos\n")
        o.append("Proyectos de investigación y creación 2014-2026 agrupados por sublínea, del más reciente al más antiguo. "
                 "† indica asignación de confianza media o baja, por revisar.\n")
        S = sorted([s for s in subs if s["línea"] == lin and pry_sub[s["nombre"]]], key=lambda s: (s["polo"] != "teórico", -len(pry_sub[s["nombre"]])))
        for s in S:
            o.append(f"### {s['nombre']} (polo {s['polo']})\n")
            for p in sorted(pry_sub[s["nombre"]], key=lambda p: -(p["año"] or 0)):
                mark = " †" if p["revisar"] == "sí" else ""
                o.append(f"- **{p['año']}** · {p['título']} — {p['investigadores']}{mark}")
            o.append("")

    o.append("## Productividad del cuerpo académico\n")
    o.append("Publicaciones y proyectos declarados por cada profesor en el levantamiento 2014-2026 (un proyecto compartido cuenta para cada participante). "
             "La columna *Líneas* indica dónde cultiva sublíneas según `08_Temas`; en negrita, los miembros del claustro recomendado.\n")
    o.append("| Profesor | Líneas | Publicaciones 2014-2026 | Publicaciones 2022-2026 | Proyectos 2014-2026 | Proyectos 2022-2026 |\n|---|---|---:|---:|---:|---:|")
    cl = {n for v in CLAUSTRO.values() for n in v}
    def lineas_de(n):
        ls = [lab for lab, lin in (("FD", L1), ("PP", L2)) if n in profs_lin[lin]]
        return " · ".join(ls) or "—"
    for n, d in sorted(prod.items(), key=lambda kv: -(kv[1]["publicaciones_2022_2026"] or 0) - (kv[1]["proyectos_2022_2026"] or 0)):
        nm = f"**{n}**" if n in cl else n
        o.append(f"| {nm} | {lineas_de(n)} | {d['publicaciones_2014_2026']} | {d['publicaciones_2022_2026']} | {d['proyectos_2014_2026']} | {d['proyectos_2022_2026']} |")
    o.append("\nFD: Fundamentos disciplinares · PP: Prácticas proyectuales.\n")

    rev = [p for p in pry if p["confianza"] == "baja"]
    o.append("## Asignaciones por revisar\n")
    o.append(f"De los {len(pry)} proyectos, {sum(p['confianza']=='alta' for p in pry)} tienen asignación de confianza alta, "
             f"{sum(p['confianza']=='media' for p in pry)} media y {len(rev)} baja. Los de confianza baja son:\n")
    for p in rev:
        o.append(f"- {p['año']} · {p['título']} ({p['investigadores']}) → {p['sublínea']}")
    o.append("\nPara corregir una asignación, cambiar la sublínea en `19_Proyectos` (menú desplegable) y regenerar este documento con "
             "`python3 tools/build_comparacion.py`. La línea y el polo de la fila deben ajustarse a la nueva sublínea.\n")

    OUT.write_text("\n".join(o))
    print(f"OK: {OUT.name} · {L1}: {A} · {L2}: {B}")


if __name__ == "__main__":
    main()
