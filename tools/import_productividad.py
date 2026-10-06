#!/usr/bin/env python3
"""Incorpora al xlsx los proyectos y la productividad del levantamiento 2014-2026.

Lee data/_import/proyectos.json (proyectos únicos, deduplicados entre
coautores y clasificados por sublínea a partir del título) y
data/_import/productividad.json (conteos por profesor), y escribe:

  19_Proyectos       un proyecto por fila, con sublínea, línea y polo derivados
  20_Productividad   publicaciones y proyectos por profesor, 2014-2026 y 2022-2026

Agrega además a 07_Investigadores (y a 08_Temas) a los profesores de planta
que figuraban en el levantamiento pero no en el mapa. Idempotente.

Uso:  python3 tools/import_productividad.py
"""
import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "mad-map-data-v2.xlsx"
IMP = ROOT / "data" / "_import"
FUENTE = "Levantamiento de productividad 2014-2026 (planilla del equipo, oct 2026)"

NUEVOS = [  # id, nombre, área (por confirmar), temas [(sublínea_id, tema)]
    ("INV-MSF", "Manuel Sanfuentes", "FCT", [("SUB-50", "Obra poética de Godofredo Iommi"), ("SUB-28", "Traducción de Amereida")]),
    ("INV-IMR", "Isabel M. Reyes", "FCT", [("SUB-28", "Traducción de Amereida")]),
    ("INV-AGR", "Alejandro Garretón", "FCT", [("SUB-43", "Tipografía y grabado en la formación del diseño gráfico"), ("SUB-51", "Referencias artísticas del campo del diseño")]),
    ("INV-LAV", "Leonardo Aravena", "FCT", [("SUB-46", "Dispositivos de interacción física"), ("SUB-19", "Aprendizaje con tecnologías de interacción")]),
    ("INV-ECR", "Erick Caro", "ECH", [("SUB-25", "Prefabricación monomaterial para vivienda económica")]),
]


def header_map(ws, row):
    return {ws.cell(row, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(row, c).value}


def last_row(ws, col=1):
    r = ws.max_row
    while r > 1 and ws.cell(r, col).value in (None, ""):
        r -= 1
    return r


def fresh_sheet(wb, name, titulo, nota, headers):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = nota
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h).font = Font(bold=True)
    return ws


def main():
    proyectos = json.loads((IMP / "proyectos.json").read_text())
    prod = json.loads((IMP / "productividad.json").read_text())
    wb = load_workbook(XLSX)

    # sublíneas: id -> (nombre, línea, polo)
    ws = wb["02_Sublineas"]; h = header_map(ws, 4)
    subs = {}
    for r in range(5, last_row(ws) + 1):
        sid = ws.cell(r, h["id"]).value
        if sid:
            subs[sid] = (ws.cell(r, h["nombre"]).value, ws.cell(r, h["línea"]).value, ws.cell(r, h["polo"]).value)

    # investigadores nuevos (idempotente)
    ws = wb["07_Investigadores"]; h = header_map(ws, 4)
    existentes = {ws.cell(r, h["id"]).value for r in range(5, last_row(ws) + 1)}
    for iid, nombre, area, _ in NUEVOS:
        if iid in existentes:
            continue
        r = last_row(ws) + 1
        ws.cell(r, h["id"], iid); ws.cell(r, h["nombre"], nombre); ws.cell(r, h["área_principal"], area)
        ws.cell(r, h["estado_perfil"], "agregado oct 2026 desde levantamiento de productividad; área y temas por confirmar")

    ws = wb["08_Temas"]; h = header_map(ws, 4)
    ya = {(ws.cell(r, 1).value, ws.cell(r, 2).value) for r in range(5, last_row(ws) + 1)}
    for _, nombre, _, temas in NUEVOS:
        for sid, tema in temas:
            if (subs[sid][0], nombre) in ya:
                continue
            r = last_row(ws) + 1
            ws.cell(r, 1, subs[sid][0]); ws.cell(r, 2, nombre); ws.cell(r, 3, tema)

    # 19_Proyectos
    ws = fresh_sheet(wb, "19_Proyectos",
        "19 · Proyectos de investigación y creación (2014-2026)",
        "Un proyecto por fila, deduplicado entre coautores. La sublínea se asignó a partir del título "
        "(confianza alta/media/baja); revisar las marcadas. Línea y polo se derivan de la sublínea. "
        "Fuente: " + FUENTE + ". No incluye a profesores externos.",
        ["id", "año", "título", "investigadores", "sublínea", "línea", "polo", "ventana_2022_2026", "confianza", "revisar", "notas"])
    for i, p in enumerate(proyectos, 5):
        nom, lin, polo = subs[p["sub"]]
        vals = [p["id"], p["anio"], p["titulo"], ", ".join(p["investigadores"]), nom, lin, polo,
                "sí" if (p["anio"] or 0) >= 2022 else "no", p["conf"], "sí" if p["conf"] != "alta" else "no", None]
        for c, v in enumerate(vals, 1):
            ws.cell(i, c, v)
    dv = DataValidation(type="list", formula1="SublineaNombres", allow_blank=True)
    dv.add(f"E5:E{5 + len(proyectos) + 300}")
    ws.add_data_validation(dv)
    for col, w in zip("ABCDEFGHIJK", [9, 6, 70, 40, 40, 26, 11, 10, 10, 8, 30]):
        ws.column_dimensions[col].width = w

    # 20_Productividad
    ws = fresh_sheet(wb, "20_Productividad",
        "20 · Productividad por profesor (2014-2026)",
        "Conteo de publicaciones y proyectos declarados por cada profesor; un proyecto compartido cuenta para "
        "cada participante. Fuente: " + FUENTE + ". Las publicaciones no distinguen indexación: para WoS/Scopus "
        "y capítulos con referato ver propuesta-dos-lineas.md.",
        ["investigador", "publicaciones_2014_2026", "publicaciones_2022_2026", "proyectos_2014_2026", "proyectos_2022_2026"])
    for i, (n, d) in enumerate(sorted(prod.items()), 5):
        for c, v in enumerate([n, d["pub"], d["pub5"], d["pry"], d["pry5"]], 1):
            ws.cell(i, c, v)
    ws.column_dimensions["A"].width = 32

    wb.save(XLSX)
    print(f"OK: {len(proyectos)} proyectos, {len(prod)} profesores, {len(NUEVOS)} investigadores nuevos")


if __name__ == "__main__":
    main()
